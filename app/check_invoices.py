# Run Invoice Check on data\<Ten>\*.xml -> result PDFs + log.csv in data\<Ten>\check\,
# then rebuild data\Check_Log.xlsx (everyone: valid or not, seller still active or not).
# Only invoices WITHOUT a check PDF yet are checked (add --all to redo everything).
# Works around Invoice Check not handling Vietnamese paths: copies XML to a temp ASCII folder, runs
# check-print.js --pdf there, then copies CheckResults\*.pdf + log.csv back.
import sys, os, csv, glob, shutil, tempfile, subprocess
try: sys.stdout.reconfigure(encoding="utf-8")
except Exception: pass
from layout import APP, DATA, LOG_COLS, people_dirs, read_check_log, check_problem

# the UI runs this with no console of its own; without this flag node would open one
NO_WINDOW = 0x08000000 if os.name == "nt" else 0

# Invoice Check is bundled: app\invoice_check\check-print.js, using app\node.exe + app\node_modules
IC      = os.path.join(APP, "invoice_check")
IC_NODE = os.path.join(APP, "node.exe")
IC_JS   = os.path.join(IC, "check-print.js")

def merge_log(person, new_rows):
    """Add/replace rows of this run in data\\<person>\\check\\log.csv (keyed by XML file name)."""
    rows = read_check_log(person)
    for r in new_rows:
        if r.get("Output"): r["Output"] = os.path.basename(r["Output"])   # temp path is deleted
        rows[r["File"]] = r
    if not rows: return
    with open(os.path.join(DATA, person, "check", "log.csv"), "w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=LOG_COLS, extrasaction="ignore"); w.writeheader()
        for k in sorted(rows): w.writerow(rows[k])

def pending(person, redo=False):
    """XML of this person that still need a check."""
    pdir = os.path.join(DATA, person); cdir = os.path.join(pdir, "check")
    xmls = sorted(glob.glob(os.path.join(pdir, "*.xml")))
    if redo: return xmls
    # already checked = in log.csv, or a check PDF whose name starts with the invoice name (renamed ones too)
    logged = set(read_check_log(person))
    have = [f.lower() for f in os.listdir(cdir)] if os.path.isdir(cdir) else []
    def done(x):
        b = os.path.basename(x); s = os.path.splitext(b)[0].lower()
        return b in logged or any(f.startswith(s) and f.endswith(".pdf") for f in have)
    return [x for x in xmls if not done(x)]

def check_all(people, redo=False):
    """ONE Invoice Check run for everybody (one browser, shared parallel pool), then split results per person.
    Temp file names get a 'NN__' prefix = person index, so equal file names of 2 people never clash."""
    todo = {p: pending(p, redo) for p in people}
    for p, xs in todo.items():
        if not xs: print(f"  {p}: nothing new to check")
    total = sum(len(xs) for xs in todo.values())
    if not total: return 0
    idx = {f"{i:02d}": p for i, p in enumerate(people)}
    tmp = tempfile.mkdtemp(prefix="ic_")
    try:
        for k, p in idx.items():
            for x in todo[p]: shutil.copy2(x, os.path.join(tmp, f"{k}__{os.path.basename(x)}"))
        node = IC_NODE if os.path.exists(IC_NODE) else "node"
        jobs = str(min(total, 5))    # check phase: 5 in parallel, 10 got throttled by the site (export is capped inside Invoice Check)
        print(f"  Checking {total} invoice(s) of {sum(1 for x in todo.values() if x)} person(s) in one run ({jobs} parallel)...")
        sys.stdout.flush()
        # stdout/stderr passed on purpose: without it the child writes to its own console
        # instead of our pipe, and the UI would show none of the Invoice Check output
        subprocess.call([node, IC_JS, tmp, "--pdf", "--jobs", jobs], cwd=IC,
                        creationflags=NO_WINDOW, stdout=sys.stdout, stderr=subprocess.STDOUT)
        res = os.path.join(tmp, "CheckResults"); n = 0
        for f in os.listdir(res) if os.path.isdir(res) else []:
            k, _, name = f.partition("__")
            if not f.lower().endswith(".pdf") or k not in idx or not name: continue
            cdir = os.path.join(DATA, idx[k], "check"); os.makedirs(cdir, exist_ok=True)
            shutil.copy2(os.path.join(res, f), os.path.join(cdir, name)); n += 1
        rows = {}
        lf = os.path.join(res, "log.csv")
        if os.path.exists(lf):
            with open(lf, encoding="utf-8-sig", newline="") as fh:
                for r in csv.DictReader(fh):
                    k, _, name = (r.get("File") or "").partition("__")
                    if k not in idx or not name: continue
                    r["File"] = name
                    if r.get("Output"): r["Output"] = os.path.basename(r["Output"]).partition("__")[2]
                    rows.setdefault(idx[k], []).append(r)
        for p, xs in todo.items():
            if not xs: continue
            os.makedirs(os.path.join(DATA, p, "check"), exist_ok=True)
            merge_log(p, rows.get(p, []))
            print(f"  {p}: {len(xs)} invoice(s) -> data\\{p}\\check")
        return n
    finally:
        shutil.rmtree(tmp, ignore_errors=True)

def write_summary():
    """data\\Check_Log.xlsx: one row per invoice for everyone in data\\, problem rows in red."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    wb = Workbook(); ws = wb.active; ws.title = "Check_Log"
    ws.append(["Person"] + LOG_COLS[:6] + ["Problem"])
    thin = Side(style="thin", color="BFBFBF"); bd = Border(thin, thin, thin, thin)
    red = PatternFill("solid", start_color="FDE2E1"); probs = []
    for p in people_dirs():
        for f, r in sorted(read_check_log(p).items()):
            pr = check_problem(r)
            ws.append([p] + [r.get(c, "") for c in LOG_COLS[:6]] + [pr])
            if pr:
                probs.append(f"{p} / {f}: {pr}")
                for c in ws[ws.max_row]: c.fill = red
    for c in ws[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=9)
        c.fill = PatternFill("solid", start_color="183665"); c.alignment = Alignment(horizontal="center")
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font = Font(name="Arial", size=9); c.border = bd
    for i, w in enumerate([20, 40, 9, 40, 14, 10, 30, 30], 1): ws.column_dimensions[chr(64+i)].width = w
    ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
    out = os.path.join(DATA, "Check_Log.xlsx")
    try: wb.save(out)
    except PermissionError: print(f"\n(!) Close {out} in Excel, then run again to refresh it."); return probs
    print(f"\nLog: {out}")
    return probs

def main():
    if not os.path.exists(IC_JS):
        sys.exit(f"Invoice Check not found: {IC_JS}")
    args = sys.argv[1:]; redo = "--all" in args
    args = [a for a in args if a != "--all"]
    people = args or people_dirs()
    if not people: sys.exit("No person folder in data\\. Make one folder per person and drop the invoices in.")
    print("Invoice Check:")
    total = check_all(people, redo)
    probs = write_summary()
    print(f"Done. {total} new check PDF.")
    if probs:
        print(f"\n{len(probs)} invoice(s) need attention (red in Check_Log.xlsx):")
        for x in probs: print("  ! " + x)
    else:
        print("All checked invoices valid, all sellers active.")

if __name__ == "__main__":
    main()
