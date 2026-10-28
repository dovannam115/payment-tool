# Local web UI for Payment Tool. Binds 127.0.0.1 only, serves app\ui and runs each step as a job.
# Started by "Payment Tool.bat" with pyembed\pythonw.exe (no console window).
#   python.exe app\server.py --diag   -> same server with the log in a console (for troubleshooting)
import os, re, sys, json, time, glob, codecs, shutil, socket, zipfile, ctypes, tempfile, threading, subprocess, webbrowser
import io, unicodedata
import urllib.request, urllib.error
from http.server import ThreadingHTTPServer, BaseHTTPRequestHandler
from urllib.parse import urlparse, parse_qs

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from layout import (APP, ROOT, DATA, DONE, PEOPLE, people_dirs, scan_full, read_check_log, check_problem,
                    done_names, pdf_text, pdf_kind)
import access as A

PORT    = 8711
HOST    = "127.0.0.1"
UI      = os.path.join(APP, "ui")
DUMP    = os.path.join(APP, "dump")
VERFILE = os.path.join(ROOT, "VERSION")
CFGFILE = os.path.join(APP, "update.json")
TOKFILE = os.path.join(DUMP, "ui_token.txt")
LOGFILE = os.path.join(DUMP, "server.log")
INPUT   = os.path.join(DATA, "Request_Input.xlsx")
CHECKLOG= os.path.join(DATA, "Check_Log.xlsx")
GUIDE   = os.path.join(ROOT, "Operating_Guide.html")
BAT     = os.path.join(ROOT, "Payment Tool.bat")
DIAG    = "--diag" in sys.argv

PY   = [sys.executable] if sys.executable else ["py", "-3.12"]
NODE = os.path.join(APP, "node.exe")
if not os.path.exists(NODE): NODE = "node"
NO_WINDOW = 0x08000000 if os.name == "nt" else 0
ANSI = re.compile(r"\x1b\[[0-9;]*[A-Za-z]")   # colour codes in child output


def log(msg):
    line = time.strftime("%H:%M:%S ") + str(msg)
    if DIAG: print(line, flush=True)
    try:
        os.makedirs(DUMP, exist_ok=True)
        with open(LOGFILE, "a", encoding="utf-8") as f: f.write(line + "\n")
    except Exception: pass


def box(msg, title="Payment Tool"):
    """Fatal error popup: pythonw has no console, so a message box is the only way to be seen."""
    if DIAG: print(msg, flush=True); return
    try: ctypes.windll.user32.MessageBoxW(0, str(msg), title, 0x10)
    except Exception: pass


def version():
    try: return open(VERFILE, encoding="utf-8").read().strip() or "0.0.0"
    except Exception: return "0.0.0"


def cfg():
    out = {"repo": "", "branch": "main", "token": ""}
    try: out.update(json.load(open(CFGFILE, encoding="utf-8")))
    except Exception: pass
    return out


# ------------------------------------------------------------------ jobs
class Job:
    """One step of the tool = a chain of child processes, output streamed to the page.
    A child that prints '[[ASK]] <text>' is waiting for the user to click a button here;
    the answer is a newline written to its stdin."""
    def __init__(self, action, steps, label):
        self.action, self.steps, self.label = action, steps, label
        self.lines, self.ask, self.rc = [], None, None
        self.progress = ""              # last \r line (a progress bar), shown on its own row
        self.state = "running"          # running | waiting | login | done
        self.started = time.time()
        self.proc = None
        self.stopped = False
        self.lock = threading.Lock()

    def put(self, text):
        with self.lock:
            self.lines.append(text); self.progress = ""

    def answer(self, text=""):
        p = self.proc
        self.ask = None
        if self.state in ("waiting", "login"): self.state = "running"
        if p and p.poll() is None:
            try:
                p.stdin.write(((text or "") + "\n").encode("utf-8")); p.stdin.flush(); return True
            except Exception: return False
        return False

    def feed(self, text, ended_with_cr):
        """One chunk of child output. A chunk that ended with \\r is a progress bar redrawing
        itself on the same line, so it replaces the progress row instead of piling up."""
        text = ANSI.sub("", text).rstrip()      # Invoice Check colours its bar, the page shows plain text
        if not text: return
        if ended_with_cr:
            with self.lock: self.progress = text
            return
        if text.startswith("[[ASK]]"):
            self.ask = text[7:].strip(); self.state = "waiting"
        elif text.startswith("[[LOGIN]]"):
            self.state = "login"; self.put(text[9:].strip())
        else:
            if self.state == "login": self.state = "running"   # user is through, work went on
            self.put(text)

    def kill_tree(self, p):
        """/T takes the children too, else the Edge window playwright opened stays behind."""
        if not p or p.poll() is not None: return
        try: subprocess.call(["taskkill", "/F", "/T", "/PID", str(p.pid)], creationflags=NO_WINDOW,
                             stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        except Exception: pass
        try: p.kill()
        except Exception: pass

    def stop(self):
        self.stopped = True
        self.kill_tree(self.proc)

    def run(self):
        try:
            for cmd, env in self.steps:
                if self.stopped: break
                e = dict(os.environ); e.update(env or {}); e["PYTHONIOENCODING"] = "utf-8"
                script = next((a for a in reversed(cmd) if a.lower().endswith((".py", ".js"))), cmd[0])
                self.put("> " + os.path.basename(script))
                # raw bytes, not text mode: text mode turns every \r into \n, and then a progress
                # bar redrawing itself would look like hundreds of separate lines.
                p = subprocess.Popen(cmd, cwd=APP, env=e, stdin=subprocess.PIPE,
                                     stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                                     bufsize=0, creationflags=NO_WINDOW)
                self.proc = p
                dec = codecs.getincrementaldecoder("utf-8")(errors="replace")
                buf, cr = "", False
                while True:
                    chunk = p.stdout.read(4096)      # unbuffered pipe: returns as soon as there is data
                    if not chunk: break
                    for ch in dec.decode(chunk):
                        if ch == "\r":
                            cr = True                # wait for the next character to know what it meant
                        elif ch == "\n":
                            self.feed(buf, False); buf = ""; cr = False      # \r\n and \n are both a line
                        else:
                            if cr:                   # \r then more text = a progress bar redrawing itself
                                self.feed(buf, True); buf = ""; cr = False
                            buf += ch
                if buf: self.feed(buf, cr)
                try:
                    rc = p.wait(timeout=25)          # stdout is closed, so this is only a safety net
                except subprocess.TimeoutExpired:
                    self.put("[step finished but did not close, closing it]")
                    self.kill_tree(p); rc = p.wait()
                self.rc = rc
                if rc != 0:
                    self.put("[exit code %d]" % rc)
                    if not self.stopped: break
        except Exception as ex:
            self.put("ERROR: %s" % ex); log("job %s failed: %r" % (self.action, ex))
        self.state = "done"; self.ask = None
        if self.stopped: self.put("[stopped]")

    def view(self, since=0):
        with self.lock: lines = self.lines[since:]
        return {"action": self.action, "label": self.label, "state": self.state,
                "ask": self.ask, "rc": self.rc, "total": len(self.lines),
                "progress": self.progress, "since": since, "lines": lines}


class Runner:
    def __init__(self): self.job = None
    @property
    def busy(self): return bool(self.job and self.job.state != "done")
    def start(self, job):
        if self.busy: return False
        self.job = job
        threading.Thread(target=job.run, daemon=True).start()
        return True


RUN = Runner()


def step_check(names, redo, tax=True):
    args = list(names) + (["--all"] if redo else []) + ([] if tax else ["--no-tax"])
    return [(PY + ["-u", os.path.join(APP, "check_invoices.py")] + args, {})]


def step_fill(names):
    env = {"TB_WEB": "1"}
    if names: env["TB_ONLY"] = ";".join(names)
    return [(PY + ["-u", os.path.join(APP, "build_input.py")], {}),
            (PY + ["-u", os.path.join(APP, "don_to_json.py")], {}),
            ([NODE, os.path.join(APP, "fill_don.js")], env)]


def step_archive(names):
    return [(PY + ["-u", os.path.join(APP, "archive_done.py")] + list(names), {})]


def step_login():
    return [([NODE, os.path.join(APP, "login.js")], {"TB_WEB": "1"})]


# ------------------------------------------------------------------ state
def person_rows():
    rows = []
    for name in people_dirs():
        sc = scan_full(name)
        invs, totrinh = sc["invoices"], sc["totrinh"]
        # an empty folder still shows: it was just made from "Thêm hồ sơ" and waits for its files
        empty = not invs and not totrinh and not sc["stray"] and not sc["misplaced"]
        logrows = read_check_log(name)
        cdir = os.path.join(DATA, name, "check")
        have = [f.lower() for f in os.listdir(cdir)] if os.path.isdir(cdir) else []
        checks = len([f for f in have if f.endswith(".pdf")])
        # same rule as check_invoices.pending(): in log.csv, or a check PDF whose name starts
        # with the invoice name (people rename them to ...-044check.pdf)
        def is_checked(iv):
            if os.path.basename(iv["xml"]) in logrows: return True
            s = iv["stem"].lower()
            return any(f.startswith(s) and f.endswith(".pdf") for f in have)

        def check_pdf(iv):
            """The check result of this invoice: same name first, else one that starts with it."""
            if iv["check"]: return iv["check"]
            s = iv["stem"].lower()
            hit = [f for f in have if f.startswith(s) and f.endswith(".pdf")]
            return os.path.join(cdir, hit[0]) if hit else ""

        def rel(p):
            return os.path.relpath(p, DATA).replace("\\", "/") if p else ""

        detail, problems, pending = [], [], 0
        for iv in invs:
            r = logrows.get(os.path.basename(iv["xml"]))
            prob = check_problem(r) if r else ""
            done = is_checked(iv)
            if not done: pending += 1
            if prob: problems.append("%s: %s" % (iv["stem"], prob))
            detail.append({"stem": iv["stem"], "so": iv["sh"], "kh": iv["kh"], "ngay": iv["ngay"],
                           "tong": iv["tong"], "seller": iv["seller"],
                           "pdf": bool(iv["pdf"]), "paired": iv["paired"],
                           "pdf_name": os.path.basename(iv["pdf"]) if iv["pdf"] else "",
                           "check": done, "problem": prob,
                           "pdf_file": rel(iv["pdf"]), "check_file": rel(check_pdf(iv))})
        rows.append({"name": name, "invoices": len(invs),
                     "pdf": sum(1 for i in invs if i["pdf"]), "checks": checks,
                     "totrinh": len(totrinh), "checked": bool(logrows),
                     "pending_check": pending, "empty": empty,
                     "problems": problems, "warnings": sc["warnings"], "detail": detail})
    return rows


# ------------------------------------------------------------------ intake (files dropped on the window)
def fold(s):
    """Lower-case, no accents: 'Đỗ Văn Nam' and 'do van  nam' compare equal."""
    s = unicodedata.normalize("NFD", s.replace("đ", "d").replace("Đ", "D"))
    return " ".join("".join(c for c in s if not unicodedata.combining(c)).lower().split())


_NAMES = {"sig": None, "list": []}
def known_names():
    """Folders already in data\\ plus names in people.xlsx (column Ten); second value = folded people.xlsx names."""
    try: sig = os.path.getmtime(PEOPLE)
    except OSError: sig = 0
    if _NAMES["sig"] != sig:
        lst = []
        if sig:
            try:
                from openpyxl import load_workbook
                ws = load_workbook(PEOPLE, read_only=True, data_only=True).active
                rows = ws.iter_rows(values_only=True)
                hdr = [str(c or "").strip() for c in next(rows, [])]
                if "Ten" in hdr:
                    i = hdr.index("Ten")
                    lst = [" ".join(str(r[i]).split()) for r in rows if r and len(r) > i and r[i]]
            except Exception as ex:
                log("people.xlsx: %s" % ex)
        _NAMES.update(sig=sig, list=lst)
    out, seen = [], set()
    for n in people_dirs() + _NAMES["list"]:
        if fold(n) not in seen: seen.add(fold(n)); out.append(n)
    return out, {fold(n) for n in _NAMES["list"]}


BAD_NAME = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
def clean_name(raw):
    """Person folder name: NFC, single spaces, snapped to a known spelling; '' if unusable."""
    n = " ".join(unicodedata.normalize("NFC", raw or "").split())
    if not n or BAD_NAME.search(n) or n.startswith(("_", ".")) or n.endswith(".") or len(n) > 80: return ""
    for k in known_names()[0]:
        if fold(k) == fold(n): return k
    return n


def zip_name(info):
    """Zip entries without the UTF-8 flag come back as cp437; most Windows zips are really UTF-8."""
    n = info.filename
    if not info.flag_bits & 0x800:
        try: n = n.encode("cp437").decode("utf-8")
        except Exception: pass
    return unicodedata.normalize("NFC", n)


def save_one(person, fname, data, sub=""):
    """Write one XML/PDF into data\\<person>\\ (or its check\\). Returns {saved|skipped, note}."""
    pdir = os.path.join(DATA, person)
    base = os.path.basename(unicodedata.normalize("NFC", fname).replace("\\", "/"))
    ext = os.path.splitext(base)[1].lower()
    if ext not in (".xml", ".pdf") or base.startswith(("~$", ".")):
        return {"skipped": base, "note": "không phải XML/PDF"}
    if base.lower() in done_names(person):
        return {"skipped": base, "note": "đã nộp trước đây (có trong hồ sơ đã hoàn tất)"}
    tdir = os.path.join(pdir, "check") if sub == "check" else pdir
    os.makedirs(tdir, exist_ok=True)
    target, note = os.path.join(tdir, base), ""
    if os.path.exists(target):
        if open(target, "rb").read() == data: return {"skipped": base, "note": "đã có"}
        b, e = os.path.splitext(base); k = 2
        while os.path.exists(os.path.join(tdir, "%s_%d%s" % (b, k, e))): k += 1
        target = os.path.join(tdir, "%s_%d%s" % (b, k, e))
        note = "trùng tên với file khác, lưu thành " + os.path.basename(target)
    with open(target, "wb") as fh: fh.write(data)
    # a check result dropped together with the invoices goes to check\ by itself
    if ext == ".pdf" and sub != "check" and pdf_kind(pdf_text(target)) == "check":
        cdir = os.path.join(pdir, "check"); os.makedirs(cdir, exist_ok=True)
        dest = os.path.join(cdir, os.path.basename(target))
        os.replace(target, dest); target = dest
        note = "là kết quả kiểm tra, đã để vào check"
    return {"saved": os.path.relpath(target, DATA).replace("\\", "/"), "note": note}


def tidy(person):
    """After an intake: rename invoice PDFs that were paired by number to their XML's name."""
    notes = []
    for iv in scan_full(person)["invoices"]:
        if iv["paired"] not in ("content", "filename"): continue
        dest = os.path.join(os.path.dirname(iv["pdf"]), iv["stem"] + ".pdf")
        if os.path.exists(dest): continue
        os.replace(iv["pdf"], dest)
        notes.append("%s: khác tên XML nhưng khớp số hóa đơn %s, đã đổi tên thành %s"
                     % (os.path.basename(iv["pdf"]), iv["sh"].lstrip("0"), os.path.basename(dest)))
    return notes


def make_dirs(raw_names):
    """Empty data\\<Họ tên>\\ folders picked in "Thêm hồ sơ" before any file is added."""
    made, bad = [], []
    for raw in raw_names:
        n = clean_name(raw)
        if not n: bad.append(raw); continue
        os.makedirs(os.path.join(DATA, n), exist_ok=True); made.append(n)
    return {"made": made, "bad": bad}


# ------------------------------------------------------------------ people list (people.xlsx, edited in the app)
PEOPLE_COLS = [("Ten", "ten"), ("Ma_CBCNV", "ma"), ("Ngan_Hang", "ngan_hang"), ("So_Tai_Khoan", "so_tk")]

def people_read():
    if not os.path.exists(PEOPLE): return {"rows": []}
    from openpyxl import load_workbook
    wb = load_workbook(PEOPLE, data_only=True)
    ws = wb["DATA"] if "DATA" in wb.sheetnames else wb.active
    hdr = [str(c.value or "").strip() for c in ws[1]]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        row = {k: (" ".join(str(r[hdr.index(col)]).split()) if col in hdr and hdr.index(col) < len(r)
                   and r[hdr.index(col)] not in (None, "") else "") for col, k in PEOPLE_COLS}
        if row["ten"]: out.append(row)
    return {"rows": out}


def people_write(rows):
    """Rewrite the DATA sheet of people.xlsx, keeping the guide sheet and the look of the file."""
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    clean, seen, dup = [], {}, []
    for r in rows:
        ten = " ".join(unicodedata.normalize("NFC", str(r.get("ten") or "")).split())
        if not ten: continue
        if BAD_NAME.search(ten): return {"error": "Tên có ký tự không hợp lệ: " + ten}
        if fold(ten) in seen: dup.append(ten); continue
        seen[fold(ten)] = 1
        clean.append({"ten": ten, **{k: " ".join(str(r.get(k) or "").split()) for _, k in PEOPLE_COLS[1:]}})
    if dup: return {"error": "Tên bị trùng: " + ", ".join(dup)}
    if os.path.exists(PEOPLE):
        wb = load_workbook(PEOPLE)
        ws = wb["DATA"] if "DATA" in wb.sheetnames else wb.active
        if ws.max_row > 1: ws.delete_rows(2, ws.max_row - 1)
        try: shutil.copy2(PEOPLE, os.path.join(DUMP, "people_prev.xlsx"))
        except Exception: pass
    else:
        wb = Workbook(); ws = wb.active; ws.title = "DATA"
    thin = Side(style="thin", color="BFBFBF"); bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    for i, (col, _) in enumerate(PEOPLE_COLS, 1):
        c = ws.cell(row=1, column=i, value=col)
        c.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", start_color="183665"); c.border = bd
        c.alignment = Alignment(horizontal="center")
    for ri, r in enumerate(clean, 2):
        for ci, (_, k) in enumerate(PEOPLE_COLS, 1):
            c = ws.cell(row=ri, column=ci, value=r[k] or None)
            c.font = Font(name="Arial", size=10); c.border = bd
            if k == "so_tk": c.number_format = "@"   # keep leading zeros of account numbers
    ws.freeze_panes = "A2"
    for k, w in zip("ABCD", (24, 16, 22, 20)): ws.column_dimensions[k].width = w
    try: wb.save(PEOPLE)
    except PermissionError:
        return {"error": "people.xlsx đang mở trong Excel. Đóng file rồi bấm Lưu lại."}
    _NAMES["sig"] = None
    return {"ok": True, "count": len(clean)}


def intake(raw_name, fname, data, sub=""):
    person = clean_name(raw_name)
    if not person: return {"error": "Tên người không hợp lệ: %s" % (raw_name or "(trống)")}
    if sub not in ("", "check"): sub = ""
    if os.path.splitext(fname)[1].lower() != ".zip":
        r = save_one(person, fname, data, sub); r["person"] = person
        return {"person": person, "files": [r]}
    out = {"person": person, "files": []}
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            for info in z.infolist():
                if info.is_dir() or "__MACOSX" in info.filename: continue
                parts = zip_name(info).replace("\\", "/").split("/")
                if os.path.splitext(parts[-1])[1].lower() not in (".xml", ".pdf"): continue
                inner = "check" if len(parts) > 1 and parts[-2].lower() == "check" else sub
                out["files"].append(save_one(person, parts[-1], z.read(info), inner))
    except zipfile.BadZipFile:
        return {"error": "File zip hỏng: %s" % fname}
    return out


def state():
    j = RUN.job
    return {"version": version(), "people": person_rows(),
            "has_input": os.path.exists(INPUT), "has_checklog": os.path.exists(CHECKLOG),
            "data_dir": DATA, "busy": RUN.busy,
            "job": ({"action": j.action, "label": j.label, "state": j.state} if j else None),
            "update": UPD.get("cache"), "access": A.STATE}


# ------------------------------------------------------------------ update from GitHub
UPD = {"cache": None, "checked": 0}
SKIP_DIRS  = {"node_modules", "pyembed", "dump", "data", "done", "dist", ".git",
              "__pycache__", ".edge-profile", ".vscode"}
KEEP_FILES = {"people.xlsx"}
ALLOW_EXT  = {".py", ".js", ".html", ".css", ".bat", ".txt", ".json", ".md", ".ico"}
ALLOW_NAME = {"VERSION"}   # no extension, but it IS the file that says which version this is


def wanted(name):
    return name in ALLOW_NAME or os.path.splitext(name)[1].lower() in ALLOW_EXT


def http_get(url, token="", timeout=25):
    req = urllib.request.Request(url, headers={"User-Agent": "PaymentTool", "Cache-Control": "no-cache"})
    if token: req.add_header("Authorization", "Bearer " + token)
    with urllib.request.urlopen(req, timeout=timeout) as r: return r.read()


def cmp_ver(a, b):
    na = [int(x) for x in str(a).replace("v", "").split(".") if x.isdigit()]
    nb = [int(x) for x in str(b).replace("v", "").split(".") if x.isdigit()]
    for i in range(max(len(na), len(nb))):
        x, y = (na[i] if i < len(na) else 0), (nb[i] if i < len(nb) else 0)
        if x != y: return 1 if x > y else -1
    return 0


def read_remote(c, name):
    """One file from the repo. The API answers with what was just pushed; raw.githubusercontent
    sits behind a CDN that keeps serving the old file for a few minutes, so it is only the fallback
    (the API allows 60 calls an hour per address)."""
    repo, branch, token = c["repo"], c["branch"] or "main", c["token"]
    try:
        api = "https://api.github.com/repos/%s/contents/%s?ref=%s" % (repo, name, branch)
        j = json.loads(http_get(api, token).decode("utf-8", "replace"))
        import base64
        return base64.b64decode(j.get("content", "")).decode("utf-8", "replace")
    except urllib.error.HTTPError as ex:
        if ex.code == 404: raise                       # repo or file really is not there
    except Exception:
        pass
    raw = "https://raw.githubusercontent.com/%s/%s/%s" % (repo, branch, name)
    return http_get(raw, token).decode("utf-8", "replace")


def update_check():
    c = cfg(); cur = version()
    if not c["repo"] or "/" not in c["repo"]:
        out = {"configured": False, "current": cur}
    else:
        try:
            latest = read_remote(c, "VERSION").strip()
            try: notes = read_remote(c, "RELEASE-NOTES.txt")
            except Exception: notes = ""
            out = {"configured": True, "current": cur, "latest": latest,
                   "available": cmp_ver(latest, cur) > 0, "notes": notes[:4000], "repo": c["repo"]}
        except urllib.error.HTTPError as ex:
            msg = ("chưa thấy repo %s trên GitHub" % c["repo"]) if ex.code in (403, 404) \
                  else "GitHub trả về lỗi %s" % ex.code
            out = {"configured": True, "current": cur, "error": msg, "repo": c["repo"]}
        except Exception as ex:
            out = {"configured": True, "current": cur,
                   "error": "không nối được GitHub (%s)" % str(ex)[:120], "repo": c["repo"]}
    UPD["cache"] = out; UPD["checked"] = time.time()
    return out


def update_apply(job):
    c = cfg()
    if not c["repo"]: job.put("update.json has no repo yet."); return
    url = "https://codeload.github.com/%s/zip/refs/heads/%s" % (c["repo"], c["branch"] or "main")
    job.put("Downloading %s ..." % url)
    data = http_get(url, c["token"], timeout=120)
    tmp = tempfile.mkdtemp(prefix="pt_upd_")
    try:
        zp = os.path.join(tmp, "src.zip")
        open(zp, "wb").write(data)
        with zipfile.ZipFile(zp) as z: z.extractall(tmp)
        roots = [os.path.join(tmp, d) for d in os.listdir(tmp) if os.path.isdir(os.path.join(tmp, d))]
        if not roots: job.put("Downloaded file has no source folder."); return
        src = roots[0]
        backup = os.path.join(ROOT, "app", "_backup_" + time.strftime("%Y%m%d_%H%M%S"))
        n = 0
        for base, dirs, files in os.walk(src):
            dirs[:] = [d for d in dirs if d not in SKIP_DIRS]
            for f in files:
                full = os.path.join(base, f)
                rel = os.path.relpath(full, src).replace("\\", "/")
                if f in KEEP_FILES or not wanted(f): continue
                dst = os.path.join(ROOT, rel.replace("/", os.sep))
                if os.path.exists(dst):
                    b = os.path.join(backup, rel.replace("/", os.sep))
                    os.makedirs(os.path.dirname(b), exist_ok=True); shutil.copy2(dst, b)
                os.makedirs(os.path.dirname(dst), exist_ok=True)
                shutil.copy2(full, dst); n += 1
                job.put("  " + rel)
        job.put("")
        job.put("%d file(s) updated. Old copies kept in %s" % (n, os.path.basename(backup) if n else "-"))
        job.put("Version on disk is now %s. Restart the tool to use it." % version())
        UPD["cache"] = None
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def update_local(zip_path, job):
    if not os.path.exists(zip_path): job.put("File not found: %s" % zip_path); return
    job.put("Installing from %s" % zip_path)
    tmp = tempfile.mkdtemp(prefix="pt_upd_")
    try:
        with zipfile.ZipFile(zip_path) as z: z.extractall(tmp)
        roots = [os.path.join(tmp, d) for d in os.listdir(tmp) if os.path.isdir(os.path.join(tmp, d))]
        src = roots[0] if len(roots) == 1 and not glob.glob(os.path.join(tmp, "*.py")) else tmp
        n = 0
        for base, dirs, files in os.walk(src):
            dirs[:] = [d for d in dirs if d not in SKIP_DIRS]
            for f in files:
                if f in KEEP_FILES or not wanted(f): continue
                full = os.path.join(base, f)
                rel = os.path.relpath(full, src)
                dst = os.path.join(ROOT, rel)
                os.makedirs(os.path.dirname(dst), exist_ok=True)
                shutil.copy2(full, dst); n += 1; job.put("  " + rel.replace("\\", "/"))
        job.put("")
        job.put("%d file(s) updated. Restart the tool to use it." % n)
        UPD["cache"] = None
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


def job_update(mode, zip_path=""):
    j = Job("update", [], "Cập nhật")
    def body():
        try:
            if mode == "local": update_local(zip_path, j)
            else: update_apply(j)
        except urllib.error.HTTPError as ex:
            j.put("GitHub returned %s. Check the repo name and branch in app\\update.json." % ex.code)
        except Exception as ex:
            j.put("Update failed: %s" % ex)
        j.state = "done"
    j.run = body
    return j


# ------------------------------------------------------------------ http
TOKEN = ""
LAST_PING = [time.time()]
SHUTDOWN = []


class Handler(BaseHTTPRequestHandler):
    server_version = "PaymentTool"
    protocol_version = "HTTP/1.1"

    def log_message(self, *a): pass

    def _send(self, code, body=b"", ctype="application/json; charset=utf-8"):
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(body)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        try: self.wfile.write(body)
        except Exception: pass

    def json(self, obj, code=200):
        self._send(code, json.dumps(obj, ensure_ascii=False).encode("utf-8"))

    def ok_token(self, q):
        t = self.headers.get("X-Token") or (q.get("t", [""])[0])
        return t == TOKEN

    def body(self):
        try: n = int(self.headers.get("Content-Length") or 0)
        except ValueError: n = 0
        if not n: return {}
        try: return json.loads(self.rfile.read(n).decode("utf-8"))
        except Exception: return {}

    def do_GET(self):
        u = urlparse(self.path); q = parse_qs(u.query); p = u.path
        if p in ("/", "/index.html"):
            f = os.path.join(UI, "index.html")
            if not os.path.exists(f): return self._send(404, b"ui missing", "text/plain")
            return self._send(200, open(f, "rb").read(), "text/html; charset=utf-8")
        if p == "/guide":
            if not os.path.exists(GUIDE): return self._send(404, b"no guide", "text/plain")
            return self._send(200, open(GUIDE, "rb").read(), "text/html; charset=utf-8")
        if p == "/api/state":
            if not self.ok_token(q): return self.json({"error": "token"}, 403)
            return self.json(state())
        if p == "/api/access":
            if not self.ok_token(q): return self.json({"error": "token"}, 403)
            if q.get("force", [""])[0] == "1": A.check("Payment Tool", version(), report=False)
            return self.json(A.status())
        if p == "/api/job":
            if not self.ok_token(q): return self.json({"error": "token"}, 403)
            j = RUN.job
            if not j: return self.json({"job": None})
            since = int((q.get("since", ["0"])[0]) or 0)
            return self.json({"job": j.view(since)})
        if p == "/api/update/check":
            if not self.ok_token(q): return self.json({"error": "token"}, 403)
            return self.json(update_check())
        if p == "/api/names":
            if not self.ok_token(q): return self.json({"error": "token"}, 403)
            names, listed = known_names()
            return self.json({"names": names, "listed": sorted(listed)})
        if p == "/api/people":
            if not self.ok_token(q): return self.json({"error": "token"}, 403)
            try: return self.json(people_read())
            except Exception as ex: return self.json({"error": "Không đọc được people.xlsx: %s" % ex}, 500)
        return self._send(404, b"not found", "text/plain")

    def do_POST(self):
        u = urlparse(self.path); q = parse_qs(u.query); p = u.path
        if not self.ok_token(q): return self.json({"error": "token"}, 403)
        if p == "/api/upload":
            # raw file bytes in the body; person / file name / subfolder in the query string
            try: n = int(self.headers.get("Content-Length") or 0)
            except ValueError: n = 0
            if n > 60 * 1024 * 1024: return self.json({"error": "File quá lớn (trên 60MB)"}, 413)
            data = self.rfile.read(n) if n else b""
            if RUN.busy: return self.json({"error": "Đang chạy một bước, chờ xong rồi thêm file"}, 409)
            g = lambda k: q.get(k, [""])[0]
            try: return self.json(intake(g("name"), g("file"), data, g("sub")))
            except Exception as ex:
                log("upload: %r" % ex); return self.json({"error": str(ex)}, 500)
        d = self.body()
        if p == "/api/people":
            try: return self.json(people_write(d.get("rows") or []))
            except Exception as ex:
                log("people save: %r" % ex); return self.json({"error": str(ex)}, 500)
        if p == "/api/mkdir":
            return self.json(make_dirs(d.get("names") or []))
        if p == "/api/tidy":
            notes = []
            for n in d.get("names") or []:
                if n in people_dirs(): notes += tidy(n)
            return self.json({"notes": notes})
        if p == "/api/ping":
            LAST_PING[0] = time.time(); return self.json({"ok": True})
        if p == "/api/run":
            action = d.get("action"); names = d.get("names") or []
            if RUN.busy: return self.json({"error": "busy"}, 409)
            if action != "update" and not A.STATE["ok"]:
                return self.json({"error": "This PC is locked. Contact %s to unlock it." % (A.STATE["contact"] or "the tool owner")}, 409)
            if action == "check":
                job = Job(action, step_check(names, bool(d.get("redo")), d.get("tax", True) is not False), "Kiểm tra hóa đơn")
            elif action == "fill":
                job = Job(action, step_fill(names), "Điền đơn lên cổng")
            elif action == "archive":
                if not names: return self.json({"error": "no names"}, 400)
                job = Job(action, step_archive(names), "Hoàn tất")
            elif action == "login":
                job = Job(action, step_login(), "Đăng nhập cổng")
            elif action == "update":
                job = job_update(d.get("mode") or "github", d.get("zip") or "")
            else:
                return self.json({"error": "unknown action"}, 400)
            RUN.start(job)
            return self.json({"ok": True})
        if p == "/api/answer":
            j = RUN.job
            return self.json({"ok": bool(j and j.answer(d.get("text") or ""))})
        if p == "/api/stop":
            j = RUN.job
            if j: j.stop()
            return self.json({"ok": True})
        if p == "/api/open":
            what = d.get("what"); target = {"input": INPUT, "checklog": CHECKLOG, "data": DATA,
                                            "done": DONE, "guide": GUIDE}.get(what)
            if what == "person": target = os.path.join(DATA, d.get("name") or "")
            if what == "file":
                # a file the page listed, always inside data\ - resolve it and check it stayed there
                want = os.path.realpath(os.path.join(DATA, d.get("path") or ""))
                root = os.path.realpath(DATA)
                target = want if (want == root or want.startswith(root + os.sep)) else None
            if not target or not os.path.exists(target):
                return self.json({"error": "Chưa có file này"}, 404)
            try: os.startfile(target)
            except Exception as ex: return self.json({"error": str(ex)}, 500)
            return self.json({"ok": True})
        if p == "/api/restart":
            restart()
            return self.json({"ok": True})
        if p == "/api/quit":
            j = RUN.job
            if j: j.stop()
            threading.Timer(0.4, lambda: SHUTDOWN and SHUTDOWN[0]()).start()
            return self.json({"ok": True})
        return self.json({"error": "not found"}, 404)


def restart():
    """Start a detached helper that waits for this process to free the port, then runs the .bat."""
    helper = ("import os,time,subprocess\n"
              "time.sleep(2.0)\n"
              "os.startfile(r'%s')\n" % BAT)
    exe = sys.executable or "pythonw.exe"
    pw = os.path.join(os.path.dirname(exe), "pythonw.exe")
    subprocess.Popen([pw if os.path.exists(pw) else exe, "-c", helper],
                     creationflags=NO_WINDOW | 0x00000008)
    threading.Timer(0.5, lambda: SHUTDOWN and SHUTDOWN[0]()).start()


def app_windows(title="Payment Tool"):
    """Visible top-level windows with this exact title (Edge --app windows carry the page <title>)."""
    if os.name != "nt": return set()
    u32 = ctypes.windll.user32
    proto = ctypes.WINFUNCTYPE(ctypes.c_bool, ctypes.c_void_p, ctypes.c_void_p)
    hits = set()
    def cb(hwnd, _):
        n = u32.GetWindowTextLengthW(hwnd)
        if n and u32.IsWindowVisible(hwnd):
            buf = ctypes.create_unicode_buffer(n + 1); u32.GetWindowTextW(hwnd, buf, n + 1)
            if buf.value == title: hits.add(hwnd)
        return True
    u32.EnumWindows(proto(cb), 0)
    return hits


def maximize_window(before, wait=20):
    """Edge ignores --start-maximized when it is already running, so maximize the NEW window ourselves
    (an old window of a restarted server may still be open for a few seconds)."""
    if os.name != "nt": return
    u32 = ctypes.windll.user32
    end = time.time() + wait
    while time.time() < end:
        new = app_windows() - before
        if new:
            for h in new:
                if not u32.IsZoomed(h): u32.ShowWindow(h, 3)   # SW_MAXIMIZE
            return
        time.sleep(0.3)


def open_ui():
    url = "http://%s:%d/?t=%s" % (HOST, PORT, TOKEN)
    if os.environ.get("PT_NO_OPEN"): log("PT_NO_OPEN set, window not opened: " + url); return
    for exe in (r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
                r"C:\Program Files\Microsoft\Edge\Application\msedge.exe"):
        if os.path.exists(exe):
            try:
                # Translate off: the page is Vietnamese and Edge would offer to translate it every time
                before = app_windows()
                subprocess.Popen([exe, "--app=" + url, "--start-maximized",
                                  "--disable-features=Translate"],
                                 creationflags=NO_WINDOW)
                t = threading.Thread(target=maximize_window, args=(before,), daemon=True)
                t.start()
                return t
            except Exception: pass
    webbrowser.open(url)


def already_running():
    s = socket.socket()
    try:
        s.connect((HOST, PORT)); return True
    except Exception:
        return False
    finally:
        s.close()


def watchdog():
    """No page talking to us for 2 minutes and nothing running -> close the server."""
    while True:
        time.sleep(10)
        if RUN.busy: LAST_PING[0] = time.time(); continue
        if time.time() - LAST_PING[0] > 120:
            log("no page for 2 minutes, closing"); SHUTDOWN and SHUTDOWN[0](); return


def main():
    global TOKEN
    os.makedirs(DUMP, exist_ok=True); os.makedirs(DATA, exist_ok=True); os.makedirs(DONE, exist_ok=True)
    if already_running():
        try: TOKEN = open(TOKFILE, encoding="utf-8").read().strip()
        except Exception: TOKEN = ""
        log("already running, opening the window")
        t = open_ui()
        if t: t.join()   # this process exits right after, so wait for the window to be maximized
        return
    TOKEN = os.urandom(16).hex()
    open(TOKFILE, "w", encoding="utf-8").write(TOKEN)
    try:
        srv = ThreadingHTTPServer((HOST, PORT), Handler)
    except OSError as ex:
        box("Không mở được cổng %d (%s).\nĐóng tool đang chạy rồi mở lại." % (PORT, ex)); return
    SHUTDOWN.append(lambda: threading.Thread(target=srv.shutdown, daemon=True).start())
    log("listening on %s:%d  v%s" % (HOST, PORT, version()))
    threading.Thread(target=watchdog, daemon=True).start()
    A.start("Payment Tool", version())
    threading.Thread(target=lambda: (time.sleep(1.5), update_check()), daemon=True).start()
    open_ui()
    try: srv.serve_forever()
    except KeyboardInterrupt: pass
    log("closed")


if __name__ == "__main__":
    try:
        main()
    except Exception as ex:
        import traceback
        log(traceback.format_exc())
        box("Tool gặp lỗi khi khởi động:\n\n%s\n\nLog: %s" % (ex, LOGFILE))
